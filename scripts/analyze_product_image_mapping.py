"""Analyze Excel catalog vs media/products images. Write JSON report; do not import."""
from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "NEW EXCEL FOR STORY OF FLOWERS.xlsx"
MEDIA = ROOT / "media" / "products"
OUT = ROOT / "data" / "import" / "product_image_mapping_report.json"

IMG_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".avif", ".gif"}

# Authoritative Excel category -> local folder (None = folder missing locally)
CAT_MAP: dict[str, str | None] = {
    "Flower Bouquet": "FLOWER BOUQUETS",
    "Anniversary Special": "ANNIVERSARY SPECIAL",
    "Anniversary special": "ANNIVERSARY SPECIAL",
    "Birthday special": "BIRTHDAY SPECIAL",
    "Bridal bouquet": "BRIDAL BOUQUETS",
    "Budget flowers (below - 100 qr)": "BUDGET FLOWERS (BELOW-100 QR)",
    "Congratulations Bouquets": "CONGRATULATIONS BOUQUETS",
    "EID AL ADA SPECIAL": "EID ALA ADHA SPECIAL",
    "EID SPECIAL SPECIAL": "EID SPECIAL ARRANGEMENTS",
    "FATHERS DAY SPECIAL": "FATHERS DAY SPECIAL",
    "Flower Bouquets Under 200 QR": "FLOWER BOUQUET UNDER 200 QR",
    "FLOWER BOUQUET WITH CHOCOLATE": "FLOWER BOUQUETS WITH CHOCOLATES",
    "FLOWER VASE ARRANGEMENT": "FLOWER VASE ARRANGEMENT",
    "FLOWER WITH PERFUMES": "FLOWER WITH PERFUMES",
    "FLOWER WITH WATCH": None,
    "FLOWER WITH GIFTS": "FLOWERS WITH GIFTS",
    "GARANGAO SPECIAL": None,
    "GET WELL SOON": "GET WELL SOON",
    "GRADUATION FLOWER NECKLACE": "GRADUATION FLOWER NECKLACE",
    "GRADUATION SPECIAL ARRANGEMENT": "GRADUATION SPECIAL ARRANGEMENTS",
    "Hajj and Umra Special Arrangement": "HAJJ AND UMRA SPECIAL ARRANGEMENT",
    "MOTHERS DAY SPECIAL ARRANGEMENT": "MOTHERS DAY SPECIAL ARRANGEMENT",
    "NEW BORN BABY FLOWER ARRANGEMENT": "NEW BORN BABY FLOWER ARRANGEMENTS",
    "PREMIUM FLOWER ARRANGEMENT": "PREMIUM  FLOWER ARRANGEMENTS",
    "Ramadan Special Flower Arrangement": "RAMADAN SPECIAL",
    "10 Roses Set": "ROSE BUNDLES",
    "20 Roses Set": "ROSE BUNDLES",
    "25 Roses Set": "ROSE BUNDLES",
    "30 ROSES SET": "ROSE BUNDLES",
    "Rose Petals": "ROSE PETALS",
    "ROSE WITH CYCLE AND TOYS": None,
    "Table Floral Centerpiece Arrangement": "TABLE ARRANGEMENTS",
    "TEACHERS DAY SPECIAL": "TEACHERS DAY SPECIAL",
    "VALENTINES DAY SPECIAL": "VALENTINES SPECIAL",
}


def norm(text: object) -> str:
    if text is None:
        return ""
    s = str(text).strip().lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    replacements = [
        ("bouqut", "bouquet"),
        ("arrangment", "arrangement"),
        ("arrangemnet", "arrangement"),
        ("fushia", "fuchsia"),
        ("fuvhia", "fuchsia"),
        ("fuchia", "fuchsia"),
        ("lilly", "lily"),
        ("pices", "pieces"),
        ("cholate", "chocolate"),
        ("surpice", "surprise"),
        ("traditinal", "traditional"),
        ("fusie", "fuchsia"),
        ("colorfull", "colorful"),
    ]
    for bad, good in replacements:
        s = s.replace(bad, good)
    s = re.sub(r"[_\-]+", " ", s)
    s = re.sub(r"[^a-z0-9\s]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def load_products() -> list[dict]:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    # Row 2 is Arabic column labels
    products: list[dict] = []
    for r in range(3, ws.max_row + 1):
        vals = {headers[i]: ws.cell(r, i + 1).value for i in range(len(headers))}
        if not vals.get("PRODUCT NAME - ENGLISH") and not vals.get("CATEGORY - ENGLISH"):
            continue
        products.append(
            {
                "row": r,
                "no": vals.get("NO"),
                "category_en": str(vals.get("CATEGORY - ENGLISH") or "").strip(),
                "category_ar": str(vals.get("CATEGORY - ARABIC") or "").strip(),
                "name_en": str(vals.get("PRODUCT NAME - ENGLISH") or "").strip(),
                "name_ar": str(vals.get("PRODUCT NAME - ARABIC") or "").strip(),
                "desc_en": str(vals.get("DESCRIPTION - ENGLISH") or "").strip(),
                "desc_ar": str(vals.get("DESCRIPTION - ARABIC") or "").strip(),
                "price": vals.get("PRICE"),
                "prep": vals.get("PREPARATION TIME"),
                "addon_en": str(vals.get("ADD-ONS ENGLISH") or "").strip(),
                "addon_ar": str(vals.get("ADD-ONS ARABIC") or "").strip(),
            }
        )
    return products


def index_images() -> tuple[list[str], dict[str, list[str]], dict[str, list[str]]]:
    all_images: list[str] = []
    by_stem: dict[str, list[str]] = defaultdict(list)
    by_folder: dict[str, list[str]] = defaultdict(list)
    for p in MEDIA.rglob("*"):
        if not p.is_file() or p.suffix.lower() not in IMG_EXTS:
            continue
        rel = p.relative_to(MEDIA).as_posix()
        top = rel.split("/")[0]
        if top.lower() == "images":
            continue
        all_images.append(rel)
        by_stem[norm(p.stem)].append(rel)
        by_folder[top].append(rel)
    return all_images, by_stem, by_folder


def find_matches(
    name_en: str,
    by_stem: dict[str, list[str]],
    pairs: list[tuple[str, str]],
) -> tuple[list[str], str]:
    n = norm(name_en)
    if not n:
        return [], "empty"
    hits = list(by_stem.get(n, []))
    kind = "exact"
    if not hits:
        kind = "contains"
        for sn, rel in pairs:
            if n and (n in sn or sn in n) and min(len(n), len(sn)) >= 10:
                hits.append(rel)
    if not hits:
        kind = "stripped"
        stripped = n
        for prefix in [
            "happy birthday special",
            "birthday special",
            "happy anniversary special",
            "anniversary special",
            "hand bouquet",
            "flower bouquet",
            "bouquet with",
            "bouquet",
        ]:
            if stripped.startswith(prefix + " "):
                stripped = stripped[len(prefix) :].strip()
                break
        if stripped and stripped != n:
            hits = list(by_stem.get(stripped, []))
            if not hits:
                for sn, rel in pairs:
                    if stripped in sn and len(stripped) >= 10:
                        hits.append(rel)
    seen: set[str] = set()
    uniq: list[str] = []
    for h in hits:
        if h not in seen:
            seen.add(h)
            uniq.append(h)
    return uniq, kind if uniq else "none"


def main() -> None:
    products = load_products()
    all_images, by_stem, by_folder = index_images()
    pairs = [(norm(Path(rel).stem), rel) for rel in all_images]

    stats: Counter[str] = Counter()
    none_by_cat: Counter[str] = Counter()
    count_by_cat: Counter[str] = Counter()
    none_samples: list[dict] = []
    multi_samples: list[dict] = []
    other_folder_samples: list[dict] = []
    used: set[str] = set()

    for p in products:
        count_by_cat[p["category_en"]] += 1
        folder = CAT_MAP.get(p["category_en"])
        matches, kind = find_matches(p["name_en"], by_stem, pairs)
        for m in matches:
            used.add(m)
        if not matches:
            stats["none"] += 1
            none_by_cat[p["category_en"]] += 1
            if len(none_samples) < 40:
                none_samples.append(
                    {"row": p["row"], "cat": p["category_en"], "name": p["name_en"]}
                )
            continue
        in_f = [m for m in matches if folder and m.startswith(folder + "/")]
        if in_f:
            stats["in_folder"] += 1
            if len(in_f) > 1:
                stats["multi_in_folder"] += 1
                if len(multi_samples) < 15:
                    multi_samples.append({"name": p["name_en"], "matches": in_f[:4]})
            if norm(Path(in_f[0]).stem) == norm(p["name_en"]):
                stats["exact_in_folder"] += 1
            else:
                stats["fuzzy_in_folder"] += 1
        else:
            stats["other_folder_only"] += 1
            if len(other_folder_samples) < 25:
                other_folder_samples.append(
                    {
                        "cat": p["category_en"],
                        "expected": folder,
                        "name": p["name_en"],
                        "got": matches[0],
                    }
                )

    cat_summary = []
    for cat, total in sorted(count_by_cat.items(), key=lambda x: -x[1]):
        folder = CAT_MAP.get(cat)
        imgs = len(by_folder.get(folder, [])) if folder else 0
        in_f = of = nf = 0
        for p in products:
            if p["category_en"] != cat:
                continue
            matches, _ = find_matches(p["name_en"], by_stem, pairs)
            preferred = [m for m in matches if folder and m.startswith(folder + "/")]
            if preferred:
                in_f += 1
            elif matches:
                of += 1
            else:
                nf += 1
        if folder is None:
            issue = "MISSING_FOLDER"
        elif in_f / total < 0.7:
            issue = "LOW_COVERAGE"
        elif abs(imgs - total) > 8:
            issue = "COUNT_MISMATCH"
        else:
            issue = "OK"
        cat_summary.append(
            {
                "excel": cat,
                "products": total,
                "folder": folder,
                "folder_images": imgs,
                "matched_in_folder": in_f,
                "matched_other_folder_only": of,
                "no_image": nf,
                "in_folder_pct": round(100 * in_f / total, 1) if total else 0,
                "issue": issue,
            }
        )

    miss: Counter[str] = Counter()
    prices: list[float] = []
    dups: Counter[str] = Counter()
    typos: list[dict] = []
    cat_ar: dict[str, set[str]] = defaultdict(set)
    typo_tokens = [
        "bouqut",
        "lilly",
        "arrangment",
        "fushia",
        "fuvhia",
        "fuchia",
        "pices",
        "traditinal",
        "fusie",
        "colorfull",
        "cholate",
        "surpice",
    ]
    for p in products:
        for f in (
            "category_en",
            "category_ar",
            "name_en",
            "name_ar",
            "desc_en",
            "desc_ar",
        ):
            if not p[f]:
                miss[f] += 1
        try:
            prices.append(float(p["price"]))
        except Exception:
            miss["price"] += 1
        dups[p["name_en"].lower()] += 1
        cat_ar[p["category_en"]].add(p["category_ar"])
        n = p["name_en"].lower()
        flags = [b for b in typo_tokens if b in n]
        if flags:
            typos.append({"row": p["row"], "name": p["name_en"], "flags": flags})

    missing_cat_hits = {}
    for cat in ("FLOWER WITH WATCH", "GARANGAO SPECIAL", "ROSE WITH CYCLE AND TOYS"):
        rows = []
        for p in products:
            if p["category_en"] != cat:
                continue
            m, _ = find_matches(p["name_en"], by_stem, pairs)
            rows.append({"name": p["name_en"], "matches_elsewhere": m[:3]})
        missing_cat_hits[cat] = rows

    junk = 0
    junk_dir = MEDIA / "images"
    if junk_dir.exists():
        junk = sum(
            1
            for p in junk_dir.rglob("*")
            if p.is_file() and p.suffix.lower() in IMG_EXTS
        )

    report = {
        "summary": {
            "excel_products": len(products),
            "catalog_images": len(all_images),
            "junk_images_folder": junk,
            **dict(stats),
            "unused_images": len(all_images) - len(used),
            "strict_in_folder_match_rate_pct": round(
                100 * stats["in_folder"] / max(len(products), 1), 1
            ),
            "any_name_match_rate_pct": round(
                100 * (len(products) - stats["none"]) / max(len(products), 1), 1
            ),
        },
        "excel_quality": {
            "missing_fields": dict(miss),
            "price_min": min(prices) if prices else None,
            "price_max": max(prices) if prices else None,
            "price_avg": round(sum(prices) / len(prices), 1) if prices else None,
            "prep": dict(Counter(str(p["prep"]) for p in products)),
            "addons_en": dict(Counter(p["addon_en"] for p in products)),
            "addons_ar": dict(Counter(p["addon_ar"] for p in products)),
            "duplicate_name_en_count": sum(1 for v in dups.values() if v > 1),
            "duplicate_samples": [
                {"name": k, "count": v}
                for k, v in sorted(dups.items(), key=lambda x: -x[1])
                if v > 1
            ][:20],
            "typo_count": len(typos),
            "typo_samples": typos[:25],
            "anniversary_case_split": {
                "Anniversary Special": count_by_cat.get("Anniversary Special", 0),
                "Anniversary special": count_by_cat.get("Anniversary special", 0),
            },
            "categories_with_multiple_ar_labels": {
                k: sorted(v) for k, v in cat_ar.items() if len(v) > 1
            },
        },
        "category_mapping_and_coverage": cat_summary,
        "naming_mismatches_excel_vs_folder": [
            {
                "excel": "EID AL ADA SPECIAL",
                "folder": "EID ALA ADHA SPECIAL",
                "note": "ADA vs ADHA",
            },
            {
                "excel": "EID SPECIAL SPECIAL",
                "folder": "EID SPECIAL ARRANGEMENTS",
                "note": "duplicate SPECIAL in excel name",
            },
            {
                "excel": "FLOWER BOUQUET WITH CHOCOLATE",
                "folder": "FLOWER BOUQUETS WITH CHOCOLATES",
            },
            {"excel": "FLOWER WITH GIFTS", "folder": "FLOWERS WITH GIFTS"},
            {"excel": "VALENTINES DAY SPECIAL", "folder": "VALENTINES SPECIAL"},
            {
                "excel": "Table Floral Centerpiece Arrangement",
                "folder": "TABLE ARRANGEMENTS",
            },
            {
                "excel": "10/20/25/30 Roses Set",
                "folder": "ROSE BUNDLES",
                "note": "4 excel categories -> 1 folder",
            },
            {
                "excel": "PREMIUM FLOWER ARRANGEMENT",
                "folder": "PREMIUM  FLOWER ARRANGEMENTS",
                "note": "double space in folder name",
            },
            {
                "excel": "Ramadan Special Flower Arrangement",
                "folder": "RAMADAN SPECIAL",
            },
            {
                "excel": "ROSE WITH CYCLE AND TOYS",
                "folder": "ROSE WITH TOYS (missing locally)",
            },
        ],
        "missing_folders": {
            "FLOWER WITH WATCH": 13,
            "GARANGAO SPECIAL": 4,
            "ROSE WITH CYCLE AND TOYS / ROSE WITH TOYS": 11,
        },
        "samples": {
            "no_match": none_samples,
            "multi_in_folder": multi_samples,
            "other_folder_only": other_folder_samples,
            "gift_excel_names": [
                p["name_en"] for p in products if p["category_en"] == "FLOWER WITH GIFTS"
            ][:12],
            "gift_image_names": [
                Path(f).stem for f in by_folder.get("FLOWERS WITH GIFTS", [])
            ][:12],
            "premium_excel_names": [
                p["name_en"]
                for p in products
                if p["category_en"] == "PREMIUM FLOWER ARRANGEMENT"
            ][:12],
            "premium_image_names": [
                Path(f).stem
                for f in by_folder.get("PREMIUM  FLOWER ARRANGEMENTS", [])
            ][:12],
            "missing_folder_product_image_search": missing_cat_hits,
        },
        "folder_image_counts": {k: len(v) for k, v in sorted(by_folder.items())},
        "blockers_needing_your_decision": [
            "Download missing image folders: FLOWER WITH WATCH (13), GARANGAO SPECIAL (4), ROSE WITH TOYS (11).",
            "Merge Anniversary Special + Anniversary special into one category?",
            "Map 10/20/25/30 Roses Set into single ROSE BUNDLES category?",
            "Fix excel typos on import (ADA->ADHA, lilly, bouqut, etc.) or keep as-is?",
            "No-image products: skip, import without photo, or wait for images?",
            "Multi-image matches: first only, all as gallery, or manual?",
            "PREPARATION TIME mostly 60 — confirm unit (minutes)?",
            "ADD-ONS almost always Greeting Card — one shared addon?",
            "Ignore media/products/images/ junk/test files?",
        ],
    }

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    print("SUMMARY")
    print(json.dumps(report["summary"], indent=2))
    print("COVERAGE_BY_CATEGORY")
    for c in sorted(cat_summary, key=lambda x: (x["in_folder_pct"], -x["products"])):
        line = (
            f"{c['in_folder_pct']:5.1f}% "
            f"in={c['matched_in_folder']:3d}/{c['products']:3d} "
            f"other={c['matched_other_folder_only']:3d} "
            f"none={c['no_image']:3d} "
            f"imgs={c['folder_images']:3d} "
            f"[{c['issue']}] "
            f"{c['excel']} => {c['folder']}"
        )
        print(line)


if __name__ == "__main__":
    main()
