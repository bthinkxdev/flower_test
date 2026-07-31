"""Story of Flowers Excel <-> media/products image mapping (shared by import tools)."""

from __future__ import annotations

import json
import re
import unicodedata
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path

from django.conf import settings
from openpyxl import load_workbook

from catalog.taxonomy import CATEGORY_GROUPS, CATEGORY_PARENT_BY_CHILD

IMG_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".avif", ".gif"}

# Excel category label -> local media/products folder (None = folder missing)
CAT_MAP: dict[str, str | None] = {
    "Flower Bouquet": "FLOWER BOUQUETS",
    "Anniversary Special": "ANNIVERSARY SPECIAL",
    "Anniversary special": "ANNIVERSARY SPECIAL",
    "Birthday special": "BIRTHDAY SPECIAL",
    "Bridal bouquet": "BRIDAL BOUQUETS",
    "Budget flowers (below - 100 qr)": "BUDGET FLOWERS (BELOW-100 QR)",
    "Congratulations Bouquets": "CONGRATULATIONS BOUQUETS",
    "EID AL ADA SPECIAL": "EID ALA ADHA SPECIAL",
    "EID SPECIAL SPECIAL": "EID SPECIAL ARRANGEMENTS",
    "FATHERS DAY SPECIAL": "FATHERS DAY SPECIAL",
    "Flower Bouquets Under 200 QR": "FLOWER BOUQUET UNDER 200 QR",
    "FLOWER BOUQUET WITH CHOCOLATE": "FLOWER BOUQUETS WITH CHOCOLATES",
    "FLOWER VASE ARRANGEMENT": "FLOWER VASE ARRANGEMENT",
    "FLOWER WITH PERFUMES": "FLOWER WITH PERFUMES",
    "FLOWER WITH WATCH": None,
    "FLOWER WITH GIFTS": "FLOWERS WITH GIFTS",
    "GARANGAO SPECIAL": None,
    "GET WELL SOON": "GET WELL SOON",
    "GRADUATION FLOWER NECKLACE": "GRADUATION FLOWER NECKLACE",
    "GRADUATION SPECIAL ARRANGEMENT": "GRADUATION SPECIAL ARRANGEMENTS",
    "Hajj and Umra Special Arrangement": "HAJJ AND UMRA SPECIAL ARRANGEMENT",
    "MOTHERS DAY SPECIAL ARRANGEMENT": "MOTHERS DAY SPECIAL ARRANGEMENT",
    "NEW BORN BABY FLOWER ARRANGEMENT": "NEW BORN BABY FLOWER ARRANGEMENTS",
    "PREMIUM FLOWER ARRANGEMENT": "PREMIUM  FLOWER ARRANGEMENTS",
    "Ramadan Special Flower Arrangement": "RAMADAN SPECIAL",
    "10 Roses Set": "ROSE BUNDLES",
    "20 Roses Set": "ROSE BUNDLES",
    "25 Roses Set": "ROSE BUNDLES",
    "30 ROSES SET": "ROSE BUNDLES",
    "Rose Petals": "ROSE PETALS",
    "ROSE WITH CYCLE AND TOYS": None,
    "Table Floral Centerpiece Arrangement": "TABLE ARRANGEMENTS",
    "TEACHERS DAY SPECIAL": "TEACHERS DAY SPECIAL",
    "VALENTINES DAY SPECIAL": "VALENTINES SPECIAL",
}

# Canonical storefront category name (EN) after merging case/spelling variants
CANONICAL_CATEGORY_EN: dict[str, str] = {
    "Anniversary Special": "Anniversary Special",
    "Anniversary special": "Anniversary Special",
    "EID AL ADA SPECIAL": "Eid Al Adha Special",
    "EID SPECIAL SPECIAL": "Eid Special Arrangements",
    "FLOWER BOUQUET WITH CHOCOLATE": "Flower Bouquets With Chocolates",
    "FLOWER WITH GIFTS": "Flowers With Gifts",
    "VALENTINES DAY SPECIAL": "Valentines Special",
    "Table Floral Centerpiece Arrangement": "Table Arrangements",
    "10 Roses Set": "Rose Bundles",
    "20 Roses Set": "Rose Bundles",
    "25 Roses Set": "Rose Bundles",
    "30 ROSES SET": "Rose Bundles",
    "PREMIUM FLOWER ARRANGEMENT": "Premium Flower Arrangements",
    "Ramadan Special Flower Arrangement": "Ramadan Special",
    "NEW BORN BABY FLOWER ARRANGEMENT": "New Born Baby Flower Arrangements",
    "GRADUATION SPECIAL ARRANGEMENT": "Graduation Special Arrangements",
    "Bridal bouquet": "Bridal Bouquets",
    "Birthday special": "Birthday Special",
    "Budget flowers (below - 100 qr)": "Budget Flowers (Below 100 QR)",
    "Flower Bouquets Under 200 QR": "Flower Bouquet Under 200 QR",
    "Flower Bouquet": "Flower Bouquets",
}

OCCASION_BY_CATEGORY: dict[str, tuple[str, str]] = {
    # slug, English name
    "Anniversary Special": ("anniversary", "Anniversary"),
    "Birthday Special": ("birthday", "Birthday"),
    "Bridal Bouquets": ("bridal", "Bridal"),
    "Congratulations Bouquets": ("congratulations", "Congratulations"),
    "Eid Al Adha Special": ("eid-al-adha", "Eid Al Adha"),
    "Eid Special Arrangements": ("eid", "Eid"),
    "FATHERS DAY SPECIAL": ("fathers-day", "Father's Day"),
    "GET WELL SOON": ("get-well-soon", "Get Well Soon"),
    "GRADUATION FLOWER NECKLACE": ("graduation", "Graduation"),
    "Graduation Special Arrangements": ("graduation", "Graduation"),
    "Hajj and Umra Special Arrangement": ("hajj-umra", "Hajj & Umra"),
    "MOTHERS DAY SPECIAL ARRANGEMENT": ("mothers-day", "Mother's Day"),
    "New Born Baby Flower Arrangements": ("new-born", "New Born"),
    "Ramadan Special": ("ramadan", "Ramadan"),
    "TEACHERS DAY SPECIAL": ("teachers-day", "Teacher's Day"),
    "Valentines Special": ("valentines", "Valentine's Day"),
    "GARANGAO SPECIAL": ("garangao", "Garangao"),
}

DEFAULT_OCCASION = ("everyday", "Everyday")

DEFAULT_XLSX = Path(settings.BASE_DIR) / "NEW EXCEL FOR STORY OF FLOWERS.xlsx"
DEFAULT_MEDIA = Path(settings.BASE_DIR) / "media" / "products"
DEFAULT_MAP_JSON = Path(settings.BASE_DIR) / "data" / "import" / "product_image_mapping_report.json"


def norm(text: object) -> str:
    if text is None:
        return ""
    s = str(text).strip().lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    replacements = [
        ("bouqut", "bouquet"),
        ("arrangment", "arrangement"),
        ("arrangemnet", "arrangement"),
        ("arrangnment", "arrangement"),
        ("fushia", "fuchsia"),
        ("fuvhia", "fuchsia"),
        ("fuchia", "fuchsia"),
        ("lilly", "lily"),
        ("pices", "pieces"),
        ("cholate", "chocolate"),
        ("surpice", "surprise"),
        ("surprice", "surprise"),
        ("traditinal", "traditional"),
        ("fusie", "fuchsia"),
        ("colorfull", "colorful"),
        ("boquet", "bouquet"),
        ("acrilic", "acrylic"),
    ]
    for bad, good in replacements:
        s = s.replace(bad, good)
    s = re.sub(r"[_\-]+", " ", s)
    s = re.sub(r"[^a-z0-9\s]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def canonical_category_en(excel_category: str) -> str:
    return CANONICAL_CATEGORY_EN.get(excel_category, excel_category)


def occasion_for_canonical_category(canonical_en: str) -> tuple[str, str]:
    if canonical_en in OCCASION_BY_CATEGORY:
        return OCCASION_BY_CATEGORY[canonical_en]
    for key, value in OCCASION_BY_CATEGORY.items():
        if key.lower() == canonical_en.lower():
            return value
    return DEFAULT_OCCASION


@dataclass
class MappedProduct:
    row: int
    no: int | None
    sku: str
    excel_category_en: str
    excel_category_ar: str
    canonical_category_en: str
    category_ar: str
    name_en: str
    name_ar: str
    desc_en: str
    desc_ar: str
    price: float | None
    preparation_minutes: int | None
    addon_en: str
    addon_ar: str
    media_folder: str | None
    image_relpaths: list[str] = field(default_factory=list)
    match_kind: str = "none"
    status: str = "no_image"
    occasion_slug: str = "everyday"
    occasion_name_en: str = "Everyday"
    notes: list[str] = field(default_factory=list)


def _parse_prep(value: object) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _parse_price(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _parse_no(value: object) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def load_excel_rows(xlsx_path: Path | None = None) -> list[dict]:
    path = xlsx_path or DEFAULT_XLSX
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows: list[dict] = []
    for r in range(3, ws.max_row + 1):
        vals = {headers[i]: ws.cell(r, i + 1).value for i in range(len(headers))}
        name_en = str(vals.get("PRODUCT NAME - ENGLISH") or "").strip()
        cat_en = str(vals.get("CATEGORY - ENGLISH") or "").strip()
        if not name_en and not cat_en:
            continue
        rows.append(
            {
                "row": r,
                "no": _parse_no(vals.get("NO")),
                "category_en": cat_en,
                "category_ar": str(vals.get("CATEGORY - ARABIC") or "").strip(),
                "name_en": name_en,
                "name_ar": str(vals.get("PRODUCT NAME - ARABIC") or "").strip(),
                "desc_en": str(vals.get("DESCRIPTION - ENGLISH") or "").strip(),
                "desc_ar": str(vals.get("DESCRIPTION - ARABIC") or "").strip(),
                "price": _parse_price(vals.get("PRICE")),
                "prep": _parse_prep(vals.get("PREPARATION TIME")),
                "addon_en": str(vals.get("ADD-ONS ENGLISH") or "").strip(),
                "addon_ar": str(vals.get("ADD-ONS ARABIC") or "").strip(),
            }
        )
    return rows


def index_media(media_root: Path | None = None) -> tuple[list[str], dict[str, list[str]]]:
    root = media_root or DEFAULT_MEDIA
    all_images: list[str] = []
    by_stem: dict[str, list[str]] = defaultdict(list)
    for p in root.rglob("*"):
        if not p.is_file() or p.suffix.lower() not in IMG_EXTS:
            continue
        rel = p.relative_to(root).as_posix()
        top = rel.split("/")[0]
        if top.lower() == "images":
            continue
        all_images.append(rel)
        by_stem[norm(p.stem)].append(rel)
    return all_images, by_stem


def find_image_matches(
    name_en: str,
    by_stem: dict[str, list[str]],
    pairs: list[tuple[str, str]],
) -> tuple[list[str], str]:
    n = norm(name_en)
    if not n:
        return [], "empty"
    hits = list(by_stem.get(n, []))
    kind = "exact"
    if not hits:
        kind = "contains"
        for sn, rel in pairs:
            if n and (n in sn or sn in n) and min(len(n), len(sn)) >= 10:
                hits.append(rel)
    if not hits:
        kind = "stripped"
        stripped = n
        for prefix in [
            "happy birthday special",
            "birthday special",
            "happy anniversary special",
            "anniversary special",
            "hand bouquet",
            "flower bouquet",
            "bouquet with",
            "bouquet",
        ]:
            if stripped.startswith(prefix + " "):
                stripped = stripped[len(prefix) :].strip()
                break
        if stripped and stripped != n:
            hits = list(by_stem.get(stripped, []))
            if not hits:
                for sn, rel in pairs:
                    if stripped in sn and len(stripped) >= 10:
                        hits.append(rel)
    seen: set[str] = set()
    uniq: list[str] = []
    for h in hits:
        if h not in seen:
            seen.add(h)
            uniq.append(h)
    return uniq, (kind if uniq else "none")


def prefer_folder(matches: list[str], folder: str | None) -> tuple[list[str], str]:
    if not matches:
        return [], "none"
    if folder:
        in_folder = [m for m in matches if m.startswith(folder + "/")]
        if in_folder:
            return in_folder, "in_folder"
        return matches, "other_folder"
    return matches, "no_folder"


def build_mapping(
    xlsx_path: Path | None = None,
    media_root: Path | None = None,
) -> dict:
    rows = load_excel_rows(xlsx_path)
    all_images, by_stem = index_media(media_root)
    pairs = [(norm(Path(rel).stem), rel) for rel in all_images]

    products: list[MappedProduct] = []
    status_counts: dict[str, int] = defaultdict(int)
    used_images: set[str] = set()

    for row in rows:
        excel_cat = row["category_en"]
        canon = canonical_category_en(excel_cat)
        folder = CAT_MAP.get(excel_cat)
        occ_slug, occ_name = occasion_for_canonical_category(canon)
        no = row["no"]
        sku = f"SOF-{no:04d}" if no is not None else f"SOF-R{row['row']}"

        matches, kind = find_image_matches(row["name_en"], by_stem, pairs)
        preferred, loc = prefer_folder(matches, folder)
        notes: list[str] = []

        if excel_cat != canon:
            notes.append(f"merged_category:{excel_cat}->{canon}")
        if folder is None:
            notes.append("missing_media_folder")
        if loc == "other_folder" and preferred:
            notes.append(f"image_outside_expected_folder:{preferred[0].split('/')[0]}")

        if preferred:
            status = "matched"
            if loc == "other_folder":
                status = "matched_other_folder"
            if len(preferred) > 1:
                status = "matched_multi"
                notes.append(f"multi_images:{len(preferred)}")
        else:
            status = "no_image"

        for img in preferred:
            used_images.add(img)

        mapped = MappedProduct(
            row=row["row"],
            no=no,
            sku=sku,
            excel_category_en=excel_cat,
            excel_category_ar=row["category_ar"],
            canonical_category_en=canon,
            category_ar=row["category_ar"],
            name_en=row["name_en"],
            name_ar=row["name_ar"],
            desc_en=row["desc_en"],
            desc_ar=row["desc_ar"],
            price=row["price"],
            preparation_minutes=row["prep"],
            addon_en=row["addon_en"],
            addon_ar=row["addon_ar"],
            media_folder=folder,
            image_relpaths=preferred[:5],
            match_kind=kind if preferred else "none",
            status=status,
            occasion_slug=occ_slug,
            occasion_name_en=occ_name,
            notes=notes,
        )
        products.append(mapped)
        status_counts[status] += 1

    categories: dict[str, dict] = {}
    for p in products:
        key = p.canonical_category_en
        if key not in categories:
            parent = CATEGORY_PARENT_BY_CHILD.get(p.canonical_category_en)
            categories[key] = {
                "name_en": p.canonical_category_en,
                "name_ar": p.category_ar,
                "parent_slug": parent.slug if parent else None,
                "parent_name_en": parent.name_en if parent else None,
                "media_folder": p.media_folder,
                "product_count": 0,
                "matched_count": 0,
            }
        categories[key]["product_count"] += 1
        if p.status.startswith("matched"):
            categories[key]["matched_count"] += 1
        # Prefer non-empty Arabic label
        if p.category_ar and not categories[key]["name_ar"]:
            categories[key]["name_ar"] = p.category_ar

    report = {
        "generated_for": "story_of_flowers_catalog_import",
        "source_xlsx": str(xlsx_path or DEFAULT_XLSX),
        "media_root": str(media_root or DEFAULT_MEDIA),
        "summary": {
            "product_count": len(products),
            "catalog_images": len(all_images),
            "unused_images": len(all_images) - len(used_images),
            "status_counts": dict(status_counts),
            "matched_any": sum(status_counts[k] for k in status_counts if k.startswith("matched")),
            "match_rate_pct": round(
                100
                * sum(status_counts[k] for k in status_counts if k.startswith("matched"))
                / max(len(products), 1),
                1,
            ),
        },
        "category_map": CAT_MAP,
        "canonical_category_map": CANONICAL_CATEGORY_EN,
        "taxonomy": [
            {
                "slug": group.slug,
                "name_en": group.name_en,
                "name_ar": group.name_ar,
                "display_order": group.display_order,
                "collections": list(group.collections),
            }
            for group in CATEGORY_GROUPS
        ],
        "categories": categories,
        "products": [asdict(p) for p in products],
    }
    return report


def write_mapping_json(
    out_path: Path | None = None,
    xlsx_path: Path | None = None,
    media_root: Path | None = None,
) -> Path:
    path = out_path or DEFAULT_MAP_JSON
    report = build_mapping(xlsx_path=xlsx_path, media_root=media_root)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def load_mapping_json(path: Path | None = None) -> dict:
    p = path or DEFAULT_MAP_JSON
    return json.loads(p.read_text(encoding="utf-8"))
